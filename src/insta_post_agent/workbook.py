from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import Story

HEADERS = {
    "News_Master": ["News ID", "Batch ID", "Region", "Headline", "Summary", "Category", "Event date", "Publisher", "Original URL", "Corroborating source URLs", "Trend score", "Trend status", "Google Trends evidence", "Score components", "Verification status", "Verification reasons", "Image path", "Caption", "Retrieval timestamp"],
    "Verification_Log": ["News ID", "Evidence URL", "Publisher", "Source independence assessment", "Claims checked", "Conflicts or uncertainty", "Decision", "Reason", "Verification timestamp"],
    "Publishing_Queue": ["Batch ID", "Target date", "Slide number", "News ID", "Image path", "Caption", "Freshness status", "Editorial approval status", "Manual publication status"],
}


def create_workbook(path: Path, stories: list[Story], selected: list[Story], batch_id: str, target_date: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "News_Master"
    for title in ("Verification_Log", "Publishing_Queue"):
        wb.create_sheet(title)
    for name, headers in HEADERS.items():
        sheet = wb[name]
        sheet.append(headers)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
        for story in stories if name == "News_Master" else []:
            sheet.append([
                story.news_id, story.batch_id, story.region, story.headline, story.summary, story.category,
                story.event_date, story.publisher, story.original_url, "\n".join(story.corroborating_urls),
                story.trend_score, story.trend_status, json.dumps(story.trend_evidence, ensure_ascii=False),
                json.dumps(story.score_components, sort_keys=True), story.verification_status,
                "; ".join(story.verification_reasons), story.image_path, story.caption, story.retrieved_at.isoformat(),
            ])
        if name == "Verification_Log":
            for story in stories:
                for evidence in story.evidence:
                    sheet.append([
                        story.news_id, evidence.url, evidence.publisher, evidence.independence,
                        "; ".join(evidence.claims_checked), "; ".join(evidence.conflicts_or_uncertainty),
                        story.verification_status, "; ".join(story.verification_reasons), story.retrieved_at.isoformat(),
                    ])
        if name == "Publishing_Queue":
            for slide_no, story in enumerate(selected, start=1):
                sheet.append([
                    batch_id, target_date, slide_no, story.news_id, story.image_path, story.caption,
                    story.freshness_status, story.verification_status, "NOT_PUBLISHED",
                ])
        _format_sheet(sheet)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    validate_workbook(path)


def _format_sheet(sheet) -> None:
    navy = "18344A"
    sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{max(1, sheet.max_row)}"
    for cell in sheet[1]:
        cell.font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 34
    widths = [18, 14, 12, 48, 60, 22, 15, 28, 58, 58, 14, 24, 58, 34, 20, 62, 40, 55, 26]
    for idx, width in enumerate(widths[: sheet.max_column], start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    thin = Side(style="hair", color="D9E1E8")
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F3F7FA")
        sheet.row_dimensions[row[0].row].height = 42


def validate_workbook(path: Path) -> bool:
    wb = load_workbook(path, data_only=True)
    try:
        if wb.sheetnames != list(HEADERS):
            raise ValueError(f"Unexpected workbook sheets: {wb.sheetnames}")
        for name, headers in HEADERS.items():
            sheet = wb[name]
            values = [sheet.cell(1, col).value for col in range(1, len(headers) + 1)]
            if values != headers or sheet.freeze_panes != "A2" or not sheet.auto_filter.ref:
                raise ValueError(f"Workbook validation failed for {name}")
        return True
    finally:
        wb.close()
